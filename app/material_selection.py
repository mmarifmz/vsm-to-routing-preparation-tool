from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Set, Tuple

from openpyxl import load_workbook
from openpyxl.styles import Alignment

from .models import norm
from .routing import ChangeRuleRow


P41_REQUIRED_SHEETS = ("MATERIAL_IN_SCOPE", "MAPL", "PLPO")
CALENDAR_FIRST_ROUND = ("WSTG004", "ECAL003", "ECAL003", "ECAL001")


def material_key(value: object) -> str:
    value = norm(value)
    return str(int(value)) if value.isdigit() else value


def split_workcenters(value: str) -> List[str]:
    return [norm(item) for item in (value or "").replace(";", ",").replace("/", ",").split(",") if norm(item)]


def workbook_sheets(path: str) -> List[str]:
    source = Path(path)
    if not source.exists():
        raise MaterialSelectionError(f"Workbook not found: {source}")
    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


@dataclass
class MaterialVariant:
    material: str
    plant: str
    routing_group: str
    counter: str
    variant_as_is: str = ""
    mrp_controller: str = ""
    production_supervisor: str = ""

    @property
    def key(self) -> Tuple[str, str, str, str]:
        return (self.material, self.plant, self.routing_group, self.counter)


@dataclass
class MaterialSelectionRow:
    material: str
    plant: str
    routing_group: str
    counter: str
    variant_as_is: str
    crid: str
    mrp_controller: str = ""
    production_supervisor: str = ""
    alternate: str = ""
    sequences: Tuple[str, str, str, str] = ("", "", "", "")
    status: str = "Draft"
    notes: str = ""

    def values(self) -> List[object]:
        return [
            self.material, self.plant, self.routing_group, self.counter,
            f"{self.routing_group}{self.counter}", self.variant_as_is,
            self.crid, self.alternate, "=Table928811[[#This Row],[Rtg Grp Counter]]", *self.sequences,
        ]


@dataclass
class MaterialSelectionSources:
    p41_path: str = ""
    code_reference_path: str = ""
    p41_sheets: Tuple[str, str, str] = P41_REQUIRED_SHEETS
    controller_sheets: Tuple[str, str] = ("", "")
    variants: Dict[Tuple[str, str, str, str], MaterialVariant] = field(default_factory=dict)
    workcenters: Dict[Tuple[str, str], Set[Tuple[str, str, str, str]]] = field(default_factory=dict)
    production_supervisors: Dict[Tuple[str, str], str] = field(default_factory=dict)
    mrp_controllers: Dict[Tuple[str, str], str] = field(default_factory=dict)
    valid_supervisor_codes: Set[Tuple[str, str]] = field(default_factory=set)
    valid_mrp_codes: Set[Tuple[str, str]] = field(default_factory=set)
    warnings: List[str] = field(default_factory=list)

    @property
    def p41_ready(self) -> bool:
        return bool(self.p41_path and self.variants and self.workcenters)

    @property
    def code_reference_ready(self) -> bool:
        return bool(self.code_reference_path and self.valid_mrp_codes and self.valid_supervisor_codes)

    @property
    def ready(self) -> bool:
        return self.p41_ready and self.code_reference_ready


class MaterialSelectionError(RuntimeError):
    pass


def _header_index(row: Iterable[object]) -> Dict[str, int]:
    return {norm(value).replace("_", " "): index for index, value in enumerate(row) if norm(value)}


def _find_header(ws, aliases: Dict[str, Set[str]], max_rows: int = 20) -> Tuple[int, Dict[str, int]]:
    best_row, best = 0, {}
    for row_no, row in enumerate(ws.iter_rows(min_row=1, max_row=min(max_rows, ws.max_row), values_only=True), 1):
        found: Dict[str, int] = {}
        for index, value in enumerate(row):
            label = norm(value).replace("_", " ")
            for field, names in aliases.items():
                if label in names:
                    found[field] = index
        if len(found) > len(best):
            best_row, best = row_no, found
    return best_row, best


def load_p41_routing(path: str, material_sheet: str = "MATERIAL_IN_SCOPE", mapl_sheet: str = "MAPL", plpo_sheet: str = "PLPO") -> MaterialSelectionSources:
    source = Path(path)
    if not source.exists():
        raise MaterialSelectionError(f"P41 Routing extraction not found: {source}")
    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        chosen = (material_sheet, mapl_sheet, plpo_sheet)
        missing = [name for name in chosen if name not in workbook.sheetnames]
        if missing:
            raise MaterialSelectionError(f"P41 workbook is missing confirmed tab(s): {', '.join(missing)}")
        material_data: Dict[Tuple[str, str], Tuple[str, str]] = {}
        ws = workbook[material_sheet]
        header = _header_index(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        required = ("MATNR", "WERKS", "DISPO", "FEVOR")
        if any(name not in header for name in required):
            raise MaterialSelectionError(f"{material_sheet} must contain MATNR, WERKS, DISPO and FEVOR.")
        for row in ws.iter_rows(min_row=2, max_col=max(header.values()) + 1, values_only=True):
            material, plant = material_key(row[header["MATNR"]]), norm(row[header["WERKS"]])
            if material and plant:
                material_data[(material, plant)] = (norm(row[header["DISPO"]]), norm(row[header["FEVOR"]]))

        variants: Dict[Tuple[str, str, str, str], MaterialVariant] = {}
        ws = workbook[mapl_sheet]
        header = _header_index(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        required = ("MATNR", "WERKS", "PLNNR", "PLNAL")
        if any(name not in header for name in required):
            raise MaterialSelectionError(f"{mapl_sheet} must contain MATNR, WERKS, PLNNR and PLNAL.")
        for row in ws.iter_rows(min_row=2, max_col=max(header.values()) + 1, values_only=True):
            material, plant = material_key(row[header["MATNR"]]), norm(row[header["WERKS"]])
            group, counter = norm(row[header["PLNNR"]]), norm(row[header["PLNAL"]]).zfill(2)
            if not all((material, plant, group, counter)):
                continue
            fallback = (norm(row[header["DISPO"]]) if "DISPO" in header else "", norm(row[header["FEVOR"]]) if "FEVOR" in header else "")
            mrp, supervisor = material_data.get((material, plant), fallback)
            item = MaterialVariant(material, plant, group, counter, mrp_controller=mrp, production_supervisor=supervisor)
            variants[item.key] = item

        operations: Dict[Tuple[str, str, str, str], List[Tuple[int, str]]] = defaultdict(list)
        by_workcenter: Dict[Tuple[str, str], Set[Tuple[str, str, str, str]]] = defaultdict(set)
        ws = workbook[plpo_sheet]
        header = _header_index(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        required = ("MATNR", "WERKS", "PLNNR", "PLNAL", "VORNR", "ARBPL")
        if any(name not in header for name in required):
            raise MaterialSelectionError(f"{plpo_sheet} must contain MATNR, WERKS, PLNNR, PLNAL, VORNR and ARBPL.")
        for row in ws.iter_rows(min_row=2, max_col=max(header.values()) + 1, values_only=True):
            key = (material_key(row[header["MATNR"]]), norm(row[header["WERKS"]]), norm(row[header["PLNNR"]]), norm(row[header["PLNAL"]]).zfill(2))
            wc = norm(row[header["ARBPL"]])
            if key not in variants or not wc:
                continue
            raw_operation = norm(row[header["VORNR"]])
            operation = int(raw_operation) if raw_operation.isdigit() else 999999
            operations[key].append((operation, wc))
            by_workcenter[(key[1], wc)].add(key)
        for key, values in operations.items():
            variants[key].variant_as_is = " - ".join(wc for _, wc in sorted(values, key=lambda item: item[0]))

        return MaterialSelectionSources(
            p41_path=str(source.resolve()), p41_sheets=chosen, variants=variants, workcenters=dict(by_workcenter),
            production_supervisors={(m, p): v[1] for (m, p), v in material_data.items()},
            mrp_controllers={(m, p): v[0] for (m, p), v in material_data.items()},
        )
    finally:
        workbook.close()


def load_code_reference(path: str, sources: MaterialSelectionSources, mrp_sheet: str, supervisor_sheet: str) -> MaterialSelectionSources:
    source = Path(path)
    if not source.exists():
        raise MaterialSelectionError(f"MRP / Production Supervisor workbook not found: {source}")
    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        missing = [name for name in (mrp_sheet, supervisor_sheet) if name not in workbook.sheetnames]
        if missing:
            raise MaterialSelectionError(f"Code-reference workbook is missing confirmed tab(s): {', '.join(missing)}")
        valid_mrp: Set[Tuple[str, str]] = set()
        valid_supervisor: Set[Tuple[str, str]] = set()
        configs = (
            (mrp_sheet, valid_mrp, {"plant": {"PLANT", "WERKS"}, "code": {"MRP CONTROLLER", "MRP CONTROLER", "DISPO"}}),
            (supervisor_sheet, valid_supervisor, {"plant": {"PLANT", "WERKS"}, "code": {"PRODN SUPERVISOR", "PRODUCTION SUPERVISOR CODE", "FEVOR"}}),
        )
        for sheet_name, target, aliases in configs:
            ws = workbook[sheet_name]
            header_row, header = _find_header(ws, aliases)
            if not {"plant", "code"}.issubset(header):
                raise MaterialSelectionError(f"{sheet_name} does not contain recognizable Plant and code columns.")
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                plant, code = norm(row[header["plant"]]), norm(row[header["code"]])
                if plant and code:
                    target.add((plant, code))
        sources.code_reference_path = str(source.resolve())
        sources.controller_sheets = (mrp_sheet, supervisor_sheet)
        sources.valid_mrp_codes = valid_mrp
        sources.valid_supervisor_codes = valid_supervisor
        return sources
    finally:
        workbook.close()


def _sequence_plan(scope_rules: Sequence[ChangeRuleRow], plant: str, crid: str) -> Tuple[Tuple[str, str, str, str], List[str]]:
    if plant == "CA02" and crid == "CALENDAR":
        return CALENDAR_FIRST_ROUND, ["CA02 Calendar first-round assumption: driver OLD WC plus FEVOR S01 and majority path; manufacturing owner review required"]
    values = [""] * 4
    notes: List[str] = []
    by_sequence: Dict[int, Set[str]] = defaultdict(set)
    for rule in scope_rules:
        if 1 <= rule.sequence <= 4 and norm(rule.new_wc):
            by_sequence[rule.sequence].add(norm(rule.new_wc))
    for sequence, candidates in by_sequence.items():
        if len(candidates) == 1:
            values[sequence - 1] = next(iter(candidates))
        elif len(candidates) > 1:
            notes.append(f"SEQ{sequence} has multiple candidate workcentres: {', '.join(sorted(candidates))}")
    return tuple(values), notes


def generate_material_selection(rows: Sequence[ChangeRuleRow], sources: MaterialSelectionSources) -> Tuple[List[MaterialSelectionRow], List[str]]:
    if not sources.ready:
        raise MaterialSelectionError("Confirm and load the P41 and code-reference sources in Step 0 first.")
    generated: Dict[Tuple[str, str, str, str, str, str], MaterialSelectionRow] = {}
    unmatched: List[str] = []
    grouped: Dict[Tuple[str, str, str], List[ChangeRuleRow]] = defaultdict(list)
    for rule in rows:
        if rule.plant and rule.crid:
            grouped[(norm(rule.plant), norm(rule.crid), rule.alternate or "")].append(rule)
    for (plant, crid, alternate), scope_rules in grouped.items():
        matched_keys: Set[Tuple[str, str, str, str]] = set()
        old_wc_frequency: Dict[str, int] = defaultdict(int)
        for rule in scope_rules:
            for old_wc in split_workcenters(rule.old_wc):
                old_wc_frequency[old_wc] += 1
        highest_frequency = max(old_wc_frequency.values(), default=0)
        # The repeated OLD WC is the material-scope driver; one-off logistics
        # workcentres remain part of the generated sequence but do not broaden scope.
        driver_old_wcs = {wc for wc, count in old_wc_frequency.items() if count == highest_frequency}
        for old_wc in driver_old_wcs:
            matched_keys.update(sources.workcenters.get((plant, old_wc), set()))
        if plant == "CA02" and crid == "CALENDAR":
            matched_keys = {key for key in matched_keys if sources.variants[key].production_supervisor == "S01"}
        if not matched_keys:
            unmatched.append(f"{plant} | {crid} | no P41 route uses driver OLD WC: {', '.join(sorted(driver_old_wcs)) or 'blank'}")
            continue
        material_scope = {(key[0], key[1]) for key in matched_keys}
        expanded_keys = {key for key in sources.variants if (key[0], key[1]) in material_scope}
        sequences, plan_notes = _sequence_plan(scope_rules, plant, crid)
        for key in expanded_keys:
            variant = sources.variants[key]
            notes = list(plan_notes)
            mrp = variant.mrp_controller or sources.mrp_controllers.get((variant.material, variant.plant), "")
            supervisor = variant.production_supervisor or sources.production_supervisors.get((variant.material, variant.plant), "")
            if not mrp:
                notes.append("Missing MRP Controller in P41 material scope")
            elif (variant.plant, mrp) not in sources.valid_mrp_codes:
                notes.append(f"MRP Controller {mrp} not found in confirmed code tab")
            if not supervisor:
                notes.append("Missing Production Supervisor in P41 material scope")
            elif (variant.plant, supervisor) not in sources.valid_supervisor_codes:
                notes.append(f"Production Supervisor {supervisor} not found in confirmed code tab")
            generated[(*key, crid, alternate)] = MaterialSelectionRow(
                material=variant.material, plant=variant.plant, routing_group=variant.routing_group,
                counter=variant.counter, variant_as_is=variant.variant_as_is, crid=crid,
                mrp_controller=mrp, production_supervisor=supervisor, alternate=alternate,
                sequences=sequences, status="Review Required" if notes else "Draft Ready", notes="; ".join(notes),
            )
    return sorted(generated.values(), key=lambda item: (item.plant, item.crid, item.material, item.routing_group, item.counter)), sorted(set(unmatched))


def export_material_selection_template(template_path: str, output_path: str, rows: Sequence[MaterialSelectionRow]) -> str:
    workbook = load_workbook(template_path)
    try:
        sheet_name = "MATR SELECTION (amin)"
        if sheet_name not in workbook.sheetnames:
            raise MaterialSelectionError(f"Template does not contain '{sheet_name}'.")
        ws = workbook[sheet_name]
        start_row, start_col = 8, 4
        for row_no in range(start_row, max(ws.max_row, start_row + len(rows)) + 1):
            for col_no in range(start_col, start_col + 13):
                ws.cell(row_no, col_no).value = None
        for row_no, item in enumerate(rows, start=start_row):
            for col_no, value in enumerate(item.values(), start=start_col):
                ws.cell(row_no, col_no).value = value
                ws.cell(row_no, col_no).alignment = Alignment(vertical="center")
        if "Table928811" in ws.tables:
            ws.tables["Table928811"].ref = f"D7:P{max(7, start_row + len(rows) - 1)}"
        workbook.save(output_path)
        return output_path
    finally:
        workbook.close()
