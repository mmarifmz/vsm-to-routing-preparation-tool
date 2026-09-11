from __future__ import annotations

import re
from collections import defaultdict, deque
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import PageRecord, ShapeRecord, norm


EXPORT_FIELDS = [
    "CRID", "PLANT", "ALTERNATE", "SEQUENCE", "SUBSEQ",
    "OLD WC", "NEW WC", "OP_DESCRIPTIONS", "RULE TYPE", "COLUMN",
]


@dataclass
class OldWorkcenterRecord:
    plant: str
    new_wc: str
    old_wc: str
    description: str = ""
    control_key: str = ""


@dataclass
class RoutingReference:
    source_path: str = ""
    records: List[OldWorkcenterRecord] = field(default_factory=list)
    _lookup: Dict[Tuple[str, str], List[OldWorkcenterRecord]] = field(default_factory=dict, init=False)

    def rebuild(self) -> None:
        self._lookup.clear()
        for row in self.records:
            self._lookup.setdefault((norm(row.plant), norm(row.new_wc)), []).append(row)

    def old_wcs(self, plant: str, new_wc: str) -> List[str]:
        values: List[str] = []
        for row in self._lookup.get((norm(plant), norm(new_wc)), []):
            for item in re.split(r"[,;/]", row.old_wc or ""):
                item = norm(item)
                if item and item not in values:
                    values.append(item)
        return values

    def description(self, plant: str, new_wc: str) -> str:
        rows = self._lookup.get((norm(plant), norm(new_wc)), [])
        return next((row.description.strip() for row in rows if row.description.strip()), "")

    def control_key(self, plant: str, new_wc: str) -> str:
        rows = self._lookup.get((norm(plant), norm(new_wc)), [])
        return next((row.control_key.strip() for row in rows if row.control_key.strip()), "")


@dataclass
class ChangeRuleRow:
    crid: str = ""
    plant: str = ""
    alternate: str = ""
    sequence: int = 0
    subseq: int = 1
    old_wc: str = ""
    new_wc: str = ""
    op_descriptions: str = ""
    rule_type: str = "ALL"
    column: str = ""
    status: str = "Ready"
    notes: str = ""
    source_shape_id: str = ""

    def export_values(self) -> List[object]:
        return [
            self.crid, self.plant, self.alternate, self.sequence, self.subseq,
            self.old_wc, self.new_wc, self.op_descriptions, self.rule_type, self.column,
        ]


class RoutingReferenceError(RuntimeError):
    pass


def _headers(ws, max_rows: int = 25) -> Tuple[int, Dict[str, int]]:
    wanted = {
        "WERKS": "plant", "PLANT": "plant",
        "ARBPL": "new_wc", "NEW WC": "new_wc", "WORK CENTER": "new_wc",
        "OLD WORK CENTER": "old_wc", "OLD WC": "old_wc",
        "STEXT": "description", "DESCRIPTION": "description",
        "STEUS": "control_key", "STEUS REF": "control_key", "CONTROL KEY": "control_key",
    }
    best: Tuple[int, Dict[str, int]] = (0, {})
    for row_no, row in enumerate(ws.iter_rows(min_row=1, max_row=min(max_rows, ws.max_row), values_only=True), 1):
        mapping: Dict[str, int] = {}
        for col_no, cell_value in enumerate(row, 1):
            value = str(cell_value or "").strip().upper().replace("_", " ")
            value = re.sub(r"\s+", " ", value)
            if value in wanted:
                mapping[wanted[value]] = col_no
        score = len(mapping)
        if score > len(best[1]):
            best = (row_no, mapping)
    return best


def load_routing_reference(path: str, sheet_name: str = "") -> RoutingReference:
    source = Path(path)
    if not source.exists():
        raise RoutingReferenceError(f"Routing reference workbook not found: {source}")
    wb = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        candidates = []
        worksheets = [wb[sheet_name]] if sheet_name and sheet_name in wb.sheetnames else ([] if sheet_name else wb.worksheets)
        if sheet_name and not worksheets:
            raise RoutingReferenceError(f"Confirmed WorkCenter tab was not found: {sheet_name}")
        for ws in worksheets:
            header_row, mapping = _headers(ws)
            if {"plant", "new_wc", "old_wc"}.issubset(mapping):
                score = len(mapping) + (5 if ws.title.upper() == "PTS03" else 0)
                candidates.append((score, ws, header_row, mapping))
        if not candidates:
            raise RoutingReferenceError(
                "No worksheet containing Plant/WERKS, New WC/ARBPL and Old WC columns was found."
            )
        _, ws, header_row, mapping = max(candidates, key=lambda x: x[0])
        result = RoutingReference(source_path=str(source.resolve()))
        for row in ws.iter_rows(min_row=header_row + 1, max_col=max(mapping.values()), values_only=True):
            plant = norm(row[mapping["plant"] - 1])
            new_wc = norm(row[mapping["new_wc"] - 1])
            old_wc = str(row[mapping["old_wc"] - 1] or "").strip()
            if not plant or not new_wc:
                continue
            description = str(row[mapping["description"] - 1] or "").strip() if mapping.get("description") else ""
            control_key = str(row[mapping["control_key"] - 1] or "").strip() if mapping.get("control_key") else ""
            result.records.append(OldWorkcenterRecord(plant, new_wc, old_wc, description, control_key))
        result.rebuild()
        return result
    finally:
        wb.close()


def clean_operation_description(text: str, workcenters: Sequence[str], fallback: str = "") -> str:
    lines = [re.sub(r"\s+", " ", line).strip(" -:\t") for line in (text or "").splitlines()]
    wcs = {norm(wc) for wc in workcenters}
    keep: List[str] = []
    role_terms = {"MILLWRIGHT", "WAREHOUSE", "QUALITY", "ENGINEERING", "PLANNER", "OPERATOR"}
    for line in lines:
        if not line:
            continue
        upper = norm(line)
        if upper in wcs:
            continue
        if upper in role_terms:
            continue
        # Remove standalone WC tokens while preserving meaningful process wording.
        cleaned = line
        for wc in wcs:
            cleaned = re.sub(rf"(?<![A-Z0-9]){re.escape(wc)}(?![A-Z0-9])", "", cleaned, flags=re.I)
        cleaned = re.sub(r"\s+", " ", cleaned).strip(" -:\t")
        if cleaned and cleaned not in keep:
            keep.append(cleaned)
    description = " - ".join(keep[:2]).strip()
    return description or fallback.strip()


def _wc_shapes(page: PageRecord) -> List[ShapeRecord]:
    # Leaf shapes are preferred; duplicate nested group text at the same visual
    # location is suppressed. The same WC at two different process steps must
    # remain two rows because SEQ/SUB are assigned within the CRID flow.
    result: List[ShapeRecord] = []
    seen: Set[Tuple[str, Tuple[str, ...], Optional[float], Optional[float]]] = set()
    for shape in page.shapes:
        if not shape.workcenters:
            continue
        if shape.has_children:
            continue
        key = (
            re.sub(r"\s+", " ", shape.text.strip()).upper(),
            tuple(sorted(shape.workcenters)),
            round(shape.pin_x, 3) if shape.pin_x is not None else None,
            round(shape.pin_y, 3) if shape.pin_y is not None else None,
        )
        if key in seen:
            continue
        seen.add(key)
        result.append(shape)
    if not result:
        result = [shape for shape in page.shapes if shape.workcenters]
    return result


def _shape_graph(page: PageRecord, candidate_ids: Set[str]) -> Tuple[Dict[str, Set[str]], Dict[str, Set[str]]]:
    outgoing: Dict[str, Set[str]] = defaultdict(set)
    incoming: Dict[str, Set[str]] = defaultdict(set)
    for conn in page.connections:
        src, dst = conn.source_shape_id, conn.target_shape_id
        if src in candidate_ids and dst in candidate_ids and src != dst:
            outgoing[src].add(dst)
            incoming[dst].add(src)
    return outgoing, incoming


def _fallback_sort_key(shape: ShapeRecord) -> Tuple[float, float, str]:
    # Visio coordinates grow upward. Most process maps flow left-to-right, then top-to-bottom.
    x = shape.pin_x if shape.pin_x is not None else 0.0
    y = shape.pin_y if shape.pin_y is not None else 0.0
    return (round(x, 3), -round(y, 3), shape.shape_id)


def _subsequence_sort_key(shape: ShapeRecord) -> Tuple[float, float, str]:
    """Order parallel operations top-to-bottom within one visual SEQ column."""
    x = shape.pin_x if shape.pin_x is not None else 0.0
    y = shape.pin_y if shape.pin_y is not None else 0.0
    return (-round(y, 3), round(x, 3), shape.shape_id)


def _geometry_layers(shapes: List[ShapeRecord], x_tolerance: float = 0.35) -> List[List[ShapeRecord]]:
    """Group visually aligned operations into department-level SEQ columns."""
    positioned = sorted(shapes, key=_fallback_sort_key)
    layers: List[List[ShapeRecord]] = []
    anchors: List[float] = []
    for shape in positioned:
        if shape.pin_x is None:
            layers.append([shape])
            anchors.append(float("inf"))
            continue
        x = float(shape.pin_x)
        if layers and anchors[-1] != float("inf") and abs(x - anchors[-1]) <= x_tolerance:
            layers[-1].append(shape)
            anchors[-1] = sum(float(item.pin_x) for item in layers[-1] if item.pin_x is not None) / len(layers[-1])
        else:
            layers.append([shape])
            anchors.append(x)
    return layers


def _point_to_shape_distance(x: float, y: float, shape: ShapeRecord) -> float:
    """Return the gap from a point to a shape's rectangular footprint."""
    if shape.pin_x is None or shape.pin_y is None:
        return float("inf")
    half_width = max(float(shape.width or 0.0) / 2.0, 0.0)
    half_height = max(float(shape.height or 0.0) / 2.0, 0.0)
    dx = max(abs(x - float(shape.pin_x)) - half_width, 0.0)
    dy = max(abs(y - float(shape.pin_y)) - half_height, 0.0)
    return (dx * dx + dy * dy) ** 0.5


def _resolved_process_edges(page: PageRecord) -> Set[Tuple[str, str]]:
    """Return directed process edges, repairing visibly attached loose ends."""
    by_id = {shape.shape_id: shape for shape in page.shapes}
    process_shapes = [shape for shape in page.shapes if shape.kind != "connector"]
    edges: Set[Tuple[str, str]] = set()

    def nearest_shape(x: Optional[float], y: Optional[float], excluded: Set[str]) -> str:
        if x is None or y is None:
            return ""
        candidates = [
            (_point_to_shape_distance(float(x), float(y), shape), shape.shape_id)
            for shape in process_shapes
            if shape.shape_id not in excluded
        ]
        if not candidates:
            return ""
        distance, shape_id = min(candidates)
        # An unglued connector endpoint normally stops on, or just beside, the
        # target box. A tight threshold avoids inventing long-distance links.
        return shape_id if distance <= 0.15 else ""

    for connection in page.connections:
        source = connection.source_shape_id if connection.source_shape_id in by_id else ""
        target = connection.target_shape_id if connection.target_shape_id in by_id else ""
        connector = by_id.get(connection.connector_id)
        if connector is not None:
            if not source:
                source = nearest_shape(connector.begin_x, connector.begin_y, {target})
            if not target:
                target = nearest_shape(connector.end_x, connector.end_y, {source})
        if source and target and source != target:
            edges.add((source, target))
    return edges


def _functional_layers(shapes: List[ShapeRecord], page: PageRecord) -> List[List[ShapeRecord]]:
    """Place WC labels on their connected process steps and follow flow order."""
    edges = _resolved_process_edges(page)
    if not edges:
        return []

    by_id = {shape.shape_id: shape for shape in page.shapes}
    nodes = {shape_id for edge in edges for shape_id in edge}
    outgoing: Dict[str, Set[str]] = defaultdict(set)
    incoming: Dict[str, Set[str]] = defaultdict(set)
    for source, target in edges:
        outgoing[source].add(target)
        incoming[target].add(source)

    component: Dict[str, int] = {}
    component_id = 0
    for start in sorted(nodes):
        if start in component:
            continue
        component_id += 1
        pending = [start]
        component[start] = component_id
        while pending:
            source = pending.pop()
            neighbours = outgoing.get(source, set()) | incoming.get(source, set())
            for target in neighbours:
                if target not in component:
                    component[target] = component_id
                    pending.append(target)

    indegree = {shape_id: len(incoming.get(shape_id, set())) for shape_id in nodes}
    current = deque(sorted(shape_id for shape_id, degree in indegree.items() if degree == 0))
    depth = {shape_id: 0 for shape_id in current}
    visited: Set[str] = set()
    while current:
        source = current.popleft()
        visited.add(source)
        for target in sorted(outgoing.get(source, set())):
            depth[target] = max(depth.get(target, 0), depth[source] + 1)
            indegree[target] -= 1
            if indegree[target] == 0:
                current.append(target)

    if len(visited) != len(nodes):
        return []

    process_nodes = [
        by_id[shape_id]
        for shape_id in nodes
        if shape_id in by_id and by_id[shape_id].kind != "connector"
    ]
    assigned: Dict[str, int] = {}
    assigned_component: Dict[str, int] = {}
    for shape in shapes:
        if shape.shape_id in depth:
            assigned[shape.shape_id] = depth[shape.shape_id]
            assigned_component[shape.shape_id] = component[shape.shape_id]
            continue

        parent_id = shape.parent_id
        while parent_id:
            if parent_id in depth:
                assigned[shape.shape_id] = depth[parent_id]
                assigned_component[shape.shape_id] = component[parent_id]
                break
            parent = by_id.get(parent_id)
            parent_id = parent.parent_id if parent is not None else ""
        if shape.shape_id in assigned or shape.pin_x is None or shape.pin_y is None:
            continue

        candidates = []
        for process in process_nodes:
            if process.pin_x is None or process.pin_y is None:
                continue
            distance = (
                (float(process.pin_x) - float(shape.pin_x)) ** 2
                + (float(process.pin_y) - float(shape.pin_y)) ** 2
            ) ** 0.5
            candidates.append((distance, process.shape_id))
        if candidates:
            distance, process_id = min(candidates)
            if distance <= 0.85:
                assigned[shape.shape_id] = depth[process_id]
                assigned_component[shape.shape_id] = component[process_id]

    # Long WC lists are commonly stacked under one connected process box. Once
    # one label is anchored, visually aligned labels inherit that same stage.
    for shape in shapes:
        if shape.shape_id in assigned or shape.pin_x is None:
            continue
        aligned = [
            other
            for other in shapes
            if other.shape_id in assigned
            and other.pin_x is not None
            and abs(float(other.pin_x) - float(shape.pin_x)) <= 0.35
        ]
        if aligned:
            nearest = min(
                aligned,
                key=lambda other: abs(float(other.pin_y or 0.0) - float(shape.pin_y or 0.0)),
            )
            assigned[shape.shape_id] = assigned[nearest.shape_id]
            assigned_component[shape.shape_id] = assigned_component[nearest.shape_id]

    # Independent diagram fragments can reuse the same graph depth even though
    # they are not the same manufacturing stage. Use connector order only when
    # every WC belongs to one coherent flow; otherwise keep the geometry fallback.
    if len(assigned) != len(shapes) or len(set(assigned_component.values())) != 1:
        return []

    layers_by_depth: Dict[int, List[ShapeRecord]] = defaultdict(list)
    for shape in shapes:
        if shape.shape_id in assigned:
            layers_by_depth[assigned[shape.shape_id]].append(shape)

    layers = [layers_by_depth[value] for value in sorted(layers_by_depth)]
    return layers


def _topological_layers(shapes: List[ShapeRecord], page: PageRecord) -> List[List[ShapeRecord]]:
    by_id = {shape.shape_id: shape for shape in shapes}
    outgoing, incoming = _shape_graph(page, set(by_id))
    if not any(outgoing.values()):
        functional = _functional_layers(shapes, page)
        if functional:
            return functional
        # When Visio connectors do not directly reference the WC leaf shapes,
        # functional flow is preferred. For pages without usable connectors,
        # aligned X positions still represent one process stage.
        return _geometry_layers(shapes)

    indegree = {shape_id: len(incoming.get(shape_id, set())) for shape_id in by_id}
    current = sorted([sid for sid, deg in indegree.items() if deg == 0], key=lambda sid: _fallback_sort_key(by_id[sid]))
    layers: List[List[ShapeRecord]] = []
    visited: Set[str] = set()
    while current:
        layer_ids = current
        layers.append([by_id[sid] for sid in layer_ids])
        next_ids: List[str] = []
        for sid in layer_ids:
            visited.add(sid)
            for dst in outgoing.get(sid, set()):
                indegree[dst] -= 1
                if indegree[dst] == 0:
                    next_ids.append(dst)
        current = sorted(set(next_ids), key=lambda sid: _fallback_sort_key(by_id[sid]))

    # Cycles/disconnected residue: append by geometry and flag later.
    residue = [shape for sid, shape in by_id.items() if sid not in visited]
    layers.extend([[shape] for shape in sorted(residue, key=_fallback_sort_key)])
    return layers


def generate_change_rule_draft(
    page: PageRecord,
    plant: str,
    crid: str,
    reference: Optional[RoutingReference] = None,
    valid_workcenters: Optional[Iterable[str]] = None,
) -> List[ChangeRuleRow]:
    shapes = _wc_shapes(page)
    layers = _topological_layers(shapes, page)
    valid_set = {norm(value) for value in (valid_workcenters or [])}
    rows: List[ChangeRuleRow] = []

    sequence = 0
    for layer in layers:
        sequence += 1
        branching = len(layer) > 1
        actual_subseq = 0
        for shape in sorted(layer, key=_subsequence_sort_key):
            for new_wc in shape.workcenters:
                actual_subseq += 1
                old_values = reference.old_wcs(plant, new_wc) if reference else []
                old_wc = old_values[0] if len(old_values) == 1 else (", ".join(old_values) if old_values else "")
                fallback_desc = reference.description(plant, new_wc) if reference else ""
                description = clean_operation_description(shape.text, shape.workcenters, fallback_desc)
                if not description and shape.parent_id:
                    by_id = {item.shape_id: item for item in page.shapes}
                    parent = by_id.get(shape.parent_id)
                    hops = 0
                    while parent is not None and hops < 3 and not description:
                        description = clean_operation_description(parent.text, shape.workcenters, fallback_desc)
                        parent = by_id.get(parent.parent_id) if parent.parent_id else None
                        hops += 1

                issues: List[str] = []
                if not crid:
                    issues.append("Missing CRID")
                if not plant:
                    issues.append("Missing Plant")
                if not new_wc:
                    issues.append("Missing NEW WC")
                if not old_wc:
                    issues.append("Missing OLD WC")
                elif len(old_values) > 1:
                    issues.append("Multiple OLD WC mappings")
                if not description:
                    issues.append("Blank operation description")
                if valid_set and norm(new_wc) not in valid_set:
                    issues.append("NEW WC not in Plant mapping")
                if branching:
                    issues.append("Branch / decision review")

                row = ChangeRuleRow(
                    crid=norm(crid),
                    plant=norm(plant),
                    alternate="",
                    sequence=sequence,
                    subseq=max(1, actual_subseq),
                    old_wc=old_wc,
                    new_wc=norm(new_wc),
                    op_descriptions=description,
                    rule_type="ALL",
                    column=f"SEQ{sequence}",
                    status="Ready" if not issues else "Review",
                    notes="; ".join(issues),
                    source_shape_id=shape.shape_id,
                )
                rows.append(row)

    # A repeated OLD WC leading to different NEW WCs is a decision/list rule,
    # not an unconditional ALL replacement.
    old_to_new: Dict[str, Set[str]] = defaultdict(set)
    for row in rows:
        for old_wc in re.split(r"[,;/]", row.old_wc or ""):
            if norm(old_wc):
                old_to_new[norm(old_wc)].add(norm(row.new_wc))
    for row in rows:
        old_values = [norm(value) for value in re.split(r"[,;/]", row.old_wc or "") if norm(value)]
        if any(len(old_to_new[value]) > 1 for value in old_values):
            row.rule_type = "LIST"

    return validate_change_rule_rows(rows, valid_set)


def validate_change_rule_rows(rows: List[ChangeRuleRow], valid_workcenters: Optional[Set[str]] = None) -> List[ChangeRuleRow]:
    valid_set = {norm(value) for value in (valid_workcenters or set())}
    grouped: Dict[Tuple[str, str], List[ChangeRuleRow]] = defaultdict(list)
    for row in rows:
        grouped[(norm(row.plant), norm(row.crid))].append(row)

    for group_rows in grouped.values():
        seen: Set[Tuple[int, int]] = set()
        sequences = sorted({row.sequence for row in group_rows if row.sequence > 0})
        gaps = set(range(1, max(sequences) + 1)).difference(sequences) if sequences else set()
        for row in group_rows:
            issues = [item.strip() for item in row.notes.split(";") if item.strip()]
            key = (row.sequence, row.subseq)
            if key in seen:
                issues.append("Duplicate SEQUENCE/SUBSEQ within CRID")
            seen.add(key)
            if not row.crid:
                issues.append("Missing CRID")
            if not row.plant:
                issues.append("Missing Plant")
            if not row.new_wc:
                issues.append("Missing NEW WC")
            if not row.old_wc:
                issues.append("Missing OLD WC")
            if not row.op_descriptions.strip():
                issues.append("Blank operation description")
            expected_column = f"SEQ{row.sequence}" if row.sequence > 0 else ""
            if row.column != expected_column:
                issues.append(f"COLUMN should be {expected_column}")
            if valid_set and row.new_wc and norm(row.new_wc) not in valid_set:
                issues.append("NEW WC not in Plant mapping")
            if gaps:
                issues.append("Sequence gap exists within CRID")
            # Deduplicate preserving order.
            unique: List[str] = []
            for issue in issues:
                if issue and issue not in unique:
                    unique.append(issue)
            row.notes = "; ".join(unique)
            row.status = "Ready" if not unique else "Review"
    return rows


def readiness_summary(rows: Sequence[ChangeRuleRow]) -> Dict[str, object]:
    total = len(rows)
    ready = sum(row.status == "Ready" for row in rows)
    review = total - ready
    missing_old = sum("Missing OLD WC" in row.notes for row in rows)
    pct = (ready / total * 100.0) if total else 0.0
    return {"total": total, "ready": ready, "review": review, "missing_old": missing_old, "readiness": pct}


def export_change_rule_subset(path: str, rows: Sequence[ChangeRuleRow]) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "CHANGE RULE DRAFT"
    ws.append(EXPORT_FIELDS)
    for row in rows:
        ws.append(row.export_values())
    header_fill = PatternFill("solid", fgColor="0F766E")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [18, 10, 12, 10, 8, 18, 18, 42, 14, 12]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return str(output)


def export_into_template(template_path: str, output_path: str, rows: Sequence[ChangeRuleRow], target_sheet: str = "CA02_ASSEMBLY") -> str:
    wb = load_workbook(template_path)
    try:
        if target_sheet not in wb.sheetnames:
            # Prefer first sheet that looks like a Change Rule template.
            candidates = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=True):
                    normalized = [str(value or "").strip().upper() for value in row]
                    if "CRID" in normalized and "NEW WC" in normalized and "RULE TYPE" in normalized:
                        candidates.append(ws.title)
                        break
            if not candidates:
                ws = wb.create_sheet(target_sheet)
                ws.append(EXPORT_FIELDS)
                header_row = 1
                start_row = 2
                col_map = {field: idx + 1 for idx, field in enumerate(EXPORT_FIELDS)}
            else:
                target_sheet = candidates[0]
        if target_sheet in wb.sheetnames:
            ws = wb[target_sheet]
            header_row = 0
            col_map: Dict[str, int] = {}
            for r in range(1, min(ws.max_row, 30) + 1):
                temp = {}
                for c in range(1, ws.max_column + 1):
                    key = str(ws.cell(r, c).value or "").strip().upper()
                    if key in EXPORT_FIELDS:
                        temp[key] = c
                if len(temp) >= 8:
                    header_row = r
                    col_map = temp
                    break
            if not header_row:
                raise RoutingReferenceError(f"Could not locate Change Rule header in sheet '{target_sheet}'.")
            start_row = header_row + 1
            # Clear existing draft rows in the required subset only.
            for r in range(start_row, max(ws.max_row, start_row) + 1):
                for field, c in col_map.items():
                    if field in EXPORT_FIELDS:
                        ws.cell(r, c).value = None

        for offset, row in enumerate(rows):
            values = dict(zip(EXPORT_FIELDS, row.export_values()))
            excel_row = start_row + offset
            for field, value in values.items():
                if field in col_map:
                    ws.cell(excel_row, col_map[field]).value = value
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output)
        return str(output)
    finally:
        wb.close()


def change_rule_row_key(row: ChangeRuleRow) -> Tuple[str, str, str, int, int, str]:
    """Stable-ish key used for rerun comparison/merge within a CRID."""
    return (
        norm(row.plant),
        norm(row.crid),
        str(row.source_shape_id or ""),
        int(row.sequence or 0),
        int(row.subseq or 1),
        norm(row.new_wc),
    )


def compare_change_rule_rows(old_rows: Sequence[ChangeRuleRow], new_rows: Sequence[ChangeRuleRow]) -> List[Dict[str, object]]:
    """Return field-level before/after changes for preview-before-rerun UI."""
    fields = [
        ("alternate", "ALTERNATE"),
        ("sequence", "SEQUENCE"),
        ("subseq", "SUBSEQ"),
        ("old_wc", "OLD WC"),
        ("new_wc", "NEW WC"),
        ("op_descriptions", "OP_DESCRIPTIONS"),
        ("rule_type", "RULE TYPE"),
        ("column", "COLUMN"),
    ]
    old_map = {change_rule_row_key(row): row for row in old_rows}
    new_map = {change_rule_row_key(row): row for row in new_rows}
    changes: List[Dict[str, object]] = []
    all_keys = list(dict.fromkeys(list(old_map) + list(new_map)))
    for key in all_keys:
        old = old_map.get(key)
        new = new_map.get(key)
        if old is None:
            changes.append({"change": "Added", "field": "ROW", "before": "", "after": f"{new.new_wc} / {new.op_descriptions}", "key": key})
            continue
        if new is None:
            changes.append({"change": "Removed", "field": "ROW", "before": f"{old.new_wc} / {old.op_descriptions}", "after": "", "key": key})
            continue
        for attr, label in fields:
            before = getattr(old, attr)
            after = getattr(new, attr)
            if str(before) != str(after):
                changes.append({"change": "Changed", "field": label, "before": before, "after": after, "key": key})
    return changes


def merge_preserve_manual_edits(old_rows: Sequence[ChangeRuleRow], new_rows: Sequence[ChangeRuleRow]) -> List[ChangeRuleRow]:
    """Merge refreshed detection with existing analyst edits.

    Matching rows preserve analyst-editable routing fields from the prior draft while
    refresh-only status/notes/source evidence comes from the newly generated row.
    Rows no longer present in the VSDX are not kept automatically.
    """
    old_by_shape: Dict[Tuple[str, str, str], List[ChangeRuleRow]] = defaultdict(list)
    old_by_wc: Dict[Tuple[str, str, str], List[ChangeRuleRow]] = defaultdict(list)
    for row in old_rows:
        old_by_shape[(norm(row.plant), norm(row.crid), str(row.source_shape_id or ""))].append(row)
        old_by_wc[(norm(row.plant), norm(row.crid), norm(row.new_wc))].append(row)

    merged: List[ChangeRuleRow] = []
    for fresh in new_rows:
        candidate = None
        shape_key = (norm(fresh.plant), norm(fresh.crid), str(fresh.source_shape_id or ""))
        wc_key = (norm(fresh.plant), norm(fresh.crid), norm(fresh.new_wc))
        if fresh.source_shape_id and old_by_shape.get(shape_key):
            candidate = old_by_shape[shape_key].pop(0)
        elif old_by_wc.get(wc_key):
            candidate = old_by_wc[wc_key].pop(0)
        if candidate is None:
            merged.append(fresh)
            continue
        merged.append(
            ChangeRuleRow(
                crid=fresh.crid,
                plant=fresh.plant,
                alternate=candidate.alternate,
                sequence=candidate.sequence,
                subseq=candidate.subseq,
                old_wc=candidate.old_wc or fresh.old_wc,
                new_wc=candidate.new_wc or fresh.new_wc,
                op_descriptions=candidate.op_descriptions or fresh.op_descriptions,
                rule_type=candidate.rule_type or fresh.rule_type,
                column=candidate.column or fresh.column,
                status=fresh.status,
                notes=fresh.notes,
                source_shape_id=fresh.source_shape_id,
            )
        )
    return merged


def department_readiness(rows: Sequence[ChangeRuleRow]) -> List[Dict[str, object]]:
    grouped: Dict[Tuple[str, str], List[ChangeRuleRow]] = defaultdict(list)
    for row in rows:
        grouped[(norm(row.plant), norm(row.crid))].append(row)
    result: List[Dict[str, object]] = []
    for (plant, crid), items in sorted(grouped.items()):
        summary = readiness_summary(items)
        result.append(
            {
                "plant": plant,
                "crid": crid,
                "rows": summary["total"],
                "ready": summary["ready"],
                "review": summary["review"],
                "missing_old": summary["missing_old"],
                "readiness": summary["readiness"],
                "status": "Ready" if summary["review"] == 0 else "Review",
            }
        )
    return result


def export_routing_generation_package(path: str, rows: Sequence[ChangeRuleRow], history: Optional[Sequence[Dict[str, object]]] = None) -> str:
    """Export a consolidated review workbook for one or many plants/CRIDs."""
    wb = Workbook()
    ws = wb.active
    ws.title = "CHANGE RULE"
    ws.append(EXPORT_FIELDS + ["STATUS", "REVIEW NOTES", "SOURCE SHAPE ID"])
    for row in rows:
        ws.append(row.export_values() + [row.status, row.notes, row.source_shape_id])

    fill = PatternFill("solid", fgColor="0F766E")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary_ws = wb.create_sheet("READINESS SUMMARY")
    summary_ws.append(["PLANT", "CRID", "ROWS", "READY", "REVIEW", "MISSING OLD WC", "READINESS %", "STATUS"])
    for item in department_readiness(rows):
        summary_ws.append([
            item["plant"], item["crid"], item["rows"], item["ready"], item["review"],
            item["missing_old"], item["readiness"] / 100.0, item["status"],
        ])
    for cell in summary_ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
    for cell in summary_ws["G"][1:]:
        cell.number_format = "0.0%"
    summary_ws.freeze_panes = "A2"

    hist_ws = wb.create_sheet("RERUN HISTORY")
    hist_ws.append(["PLANT", "CRID", "GENERATED AT", "RERUN AT", "SOURCE VSDX", "STATUS", "COUNT"])
    for item in history or []:
        hist_ws.append([
            item.get("plant", ""), item.get("crid", ""), item.get("generated_at", ""),
            item.get("rerun_at", ""), item.get("source", ""), item.get("status", ""), item.get("count", 0),
        ])
    for cell in hist_ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)

    for sheet in wb.worksheets:
        for col_cells in sheet.columns:
            max_len = max((len(str(cell.value or "")) for cell in col_cells), default=8)
            sheet.column_dimensions[col_cells[0].column_letter].width = min(max(max_len + 2, 10), 48)
        for row_cells in sheet.iter_rows():
            for cell in row_cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return str(output)
