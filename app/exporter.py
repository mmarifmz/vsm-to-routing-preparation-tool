from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Dict, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import FileRecord, MappingIndex


HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def safe(value):
    if value is None:
        return ""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def write_table(ws, headers, rows):
    ws.append(list(headers))
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = BORDER

    for row in rows:
        ws.append([safe(value) for value in row])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    widths: Dict[int, int] = {}
    for row in ws.iter_rows(
        min_row=1, max_row=min(ws.max_row, 2500)
    ):
        for cell in row:
            widths[cell.column] = min(
                55,
                max(
                    widths.get(cell.column, 0),
                    len(str(cell.value or "")) + 2,
                ),
            )
            cell.alignment = Alignment(
                vertical="top", wrap_text=True
            )
            cell.border = BORDER

    for index, width in widths.items():
        ws.column_dimensions[get_column_letter(index)].width = max(
            10, width
        )


def export(
    output,
    files: Sequence[FileRecord],
    mapping: MappingIndex | None,
    logs,
):
    workbook = Workbook()
    workbook.remove(workbook.active)

    page_summary = workbook.create_sheet("Page_Summary")
    rows = []
    for file in files:
        for page in file.pages:
            rows.append(
                [
                    file.file_name,
                    file.path,
                    file.plant,
                    page.name,
                    page.name_u,
                    page.page_id,
                    page.width,
                    page.height,
                    len(page.shapes),
                    page.text_shape_count,
                    page.connector_count,
                    len(page.connections),
                    len(page.workcenter_counts),
                    "; ".join(page.workcenter_counts),
                    "; ".join(page.matched_wcs),
                    "; ".join(page.unmatched_wcs),
                    "; ".join(
                        f"{department}:{count}"
                        for department, count
                        in page.candidate_counts.items()
                    ),
                    page.suggested_department,
                    page.confidence,
                    page.assigned_department,
                    page.assignment_mode,
                    file.status,
                    file.error,
                ]
            )
    write_table(
        page_summary,
        [
            "FILE",
            "FULL PATH",
            "PLANT",
            "TAB NAME",
            "TAB NAMEU",
            "PAGE ID",
            "PAGE WIDTH",
            "PAGE HEIGHT",
            "SHAPE COUNT",
            "TEXT SHAPE COUNT",
            "CONNECTOR SHAPE COUNT",
            "CONNECTION COUNT",
            "DISTINCT WORKCENTER COUNT",
            "WORKCENTERS",
            "MAPPED WORKCENTERS",
            "UNMATCHED WORKCENTERS",
            "DEPARTMENT / CRID CANDIDATES",
            "SUGGESTED DEPARTMENT / CRID",
            "CONFIDENCE",
            "ASSIGNED DEPARTMENT / CRID",
            "ASSIGNMENT MODE",
            "STATUS",
            "ERROR",
        ],
        rows,
    )
    for cell in page_summary["S"][1:]:
        cell.number_format = "0.0%"

    wc_detail = workbook.create_sheet("Workcenter_Detail")
    rows = []
    for file in files:
        for page in file.pages:
            for shape in page.shapes:
                for workcenter in shape.workcenters:
                    departments = (
                        mapping.departments_for_wc(
                            file.plant, workcenter
                        )
                        if mapping
                        else []
                    )
                    rows.append(
                        [
                            file.file_name,
                            file.plant,
                            page.name,
                            page.assigned_department,
                            shape.shape_id,
                            shape.name,
                            shape.text,
                            workcenter,
                            "; ".join(departments),
                            "Y" if departments else "N",
                            (
                                "Y"
                                if page.assigned_department
                                and page.assigned_department
                                in departments
                                else "N"
                            ),
                        ]
                    )
    write_table(
        wc_detail,
        [
            "FILE",
            "PLANT",
            "TAB NAME",
            "ASSIGNED DEPARTMENT / CRID",
            "SHAPE ID",
            "SHAPE NAME",
            "SHAPE TEXT",
            "WORKCENTER",
            "MAPPED DEPARTMENTS / CRIDS",
            "MAPPING FOUND",
            "MATCHES ASSIGNED",
        ],
        rows,
    )

    shape_metadata = workbook.create_sheet("Shape_Metadata")
    rows = []
    for file in files:
        for page in file.pages:
            for shape in page.shapes:
                rows.append(
                    [
                        file.file_name,
                        file.plant,
                        page.name,
                        page.assigned_department,
                        shape.shape_id,
                        shape.parent_id,
                        shape.name,
                        shape.name_u,
                        shape.shape_type,
                        shape.kind,
                        shape.master_id,
                        shape.text,
                        shape.pin_x,
                        shape.pin_y,
                        shape.width,
                        shape.height,
                        shape.begin_x,
                        shape.begin_y,
                        shape.end_x,
                        shape.end_y,
                        shape.angle,
                        "; ".join(shape.workcenters),
                        "; ".join(
                            f"{key}={value}"
                            for key, value
                            in shape.properties.items()
                        ),
                    ]
                )
    write_table(
        shape_metadata,
        [
            "FILE",
            "PLANT",
            "TAB NAME",
            "ASSIGNED DEPARTMENT / CRID",
            "SHAPE ID",
            "PARENT SHAPE ID",
            "SHAPE NAME",
            "SHAPE NAMEU",
            "SHAPE TYPE",
            "KIND",
            "MASTER ID",
            "TEXT",
            "PIN X",
            "PIN Y",
            "WIDTH",
            "HEIGHT",
            "BEGIN X",
            "BEGIN Y",
            "END X",
            "END Y",
            "ANGLE",
            "DETECTED WORKCENTERS",
            "CUSTOM PROPERTIES",
        ],
        rows,
    )

    connections = workbook.create_sheet("Connections")
    rows = []
    for file in files:
        for page in file.pages:
            shapes = {
                shape.shape_id: shape for shape in page.shapes
            }
            for connection in page.connections:
                source = shapes.get(
                    connection.source_shape_id
                )
                target = shapes.get(
                    connection.target_shape_id
                )
                rows.append(
                    [
                        file.file_name,
                        file.plant,
                        page.name,
                        connection.connector_id,
                        connection.source_shape_id,
                        source.text if source else "",
                        connection.target_shape_id,
                        target.text if target else "",
                        connection.source_cell,
                        connection.target_cell,
                    ]
                )
    write_table(
        connections,
        [
            "FILE",
            "PLANT",
            "TAB NAME",
            "CONNECTOR ID",
            "SOURCE SHAPE ID",
            "SOURCE TEXT",
            "TARGET SHAPE ID",
            "TARGET TEXT",
            "SOURCE CELL",
            "TARGET CELL",
        ],
        rows,
    )

    department_plant = workbook.create_sheet(
        "Department_Plant_List"
    )
    rows = []
    seen_pairs = set()
    if mapping:
        for record in mapping.records:
            pair = (record.department, record.plant)
            if pair in seen_pairs:
                continue
            seen_pairs.add(pair)
            rows.append(
                [
                    record.department,
                    record.plant,
                    "; ".join(
                        mapping.sources_for_department(
                            record.plant,
                            record.department,
                        )
                    ),
                ]
            )
    write_table(
        department_plant,
        [
            "DEPARTMENT / CRID",
            "PLANT",
            "SOURCE",
        ],
        sorted(rows, key=lambda row: (row[1], row[0])),
    )

    snapshot = workbook.create_sheet("Mapping_Snapshot")
    rows = []
    if mapping:
        for record in mapping.records:
            rows.append(
                [
                    record.department,
                    record.plant,
                    record.workcenter,
                    record.source_file,
                    record.source_sheet,
                    record.source_row,
                ]
            )
    write_table(
        snapshot,
        [
            "DEPARTMENT / CRID",
            "PLANT",
            "WORKCENTER",
            "SOURCE FILE",
            "SOURCE SHEET / TYPE",
            "SOURCE ROW",
        ],
        rows,
    )

    run_log = workbook.create_sheet("Run_Log")
    write_table(
        run_log,
        ["EXPORTED AT", "LOG"],
        [
            [
                datetime.now().isoformat(timespec="seconds"),
                line,
            ]
            for line in logs
        ],
    )

    about = workbook.create_sheet("About", 0)
    about.sheet_view.showGridLines = False
    about["A1"] = "VSM to ROUTING Preparation Tool v1.3.8"
    about["A1"].font = Font(
        size=18, bold=True, color="0F766E"
    )
    about["A3"] = "Generated"
    about["B3"] = datetime.now()
    about["B3"].number_format = "yyyy-mm-dd hh:mm:ss"
    about["A4"] = "VSDX files"
    about["B4"] = len(files)
    about["A5"] = "Pages"
    about["B5"] = sum(len(file.pages) for file in files)
    about["A6"] = "Mapping sources"
    about["B6"] = (
        mapping.source_path if mapping else "Not loaded"
    )
    about["A7"] = "Department-Plant pairs"
    about["B7"] = (
        mapping.department_plant_pair_count
        if mapping
        else 0
    )
    about["A8"] = "Workcenter mapping records"
    about["B8"] = (
        mapping.workcenter_record_count
        if mapping
        else 0
    )
    about.column_dimensions["A"].width = 28
    about.column_dimensions["B"].width = 90

    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return str(output_path)
